"""Bulk import of a research cohort (OASIS-1) into patients and scans.

Two input shapes:

* a CSV of already-derived biomarkers, one row per session, with columns
  `subject_id,cdr,age,sex,icv_mm3,hipp_mm3,hipp_norm,vbr,cortex_gm_norm`
  and an optional mean cortical thickness column;
* a folder of per-subject FreeSurfer `stats/` directories plus the OASIS
  demographics spreadsheet (`oasis_cross-sectional*.xlsx` or `.csv`).

Rules that hold for both:

* **A missing CDR is never turned into a number.** A row without a label is
  skipped and counted. An earlier script used `fillna(0.0)` here and so
  labelled every unassessed subject "healthy"; this module has no code
  path that supplies a default label.
* `_MR2` sessions are same-subject reliability rescans. They are skipped
  unless asked for, so a subject is not counted twice.
* The import is idempotent on `external_subject_id`.
* Every scan goes through `scan_service.process_derived_biomarkers`, the
  same path the API endpoints use, so QC flags and the calibration gate
  apply exactly as they do to a single upload.

The CDR is stored as the scan's research-only `reference_label` and is not
exposed by any triage endpoint.
"""

import csv
import datetime
import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.identifiers import validate_external_subject_id
from app.core.logging import get_logger
from app.models.enums import DataSource, ExtractionMethod, Sex
from app.repositories import patient_repo
from app.services import scan_service
from app.services.biomarker_derivation import (
    DerivedBiomarkers,
    derive_from_stats,
    derive_from_values,
)
from app.services.calibration import data_domain_for
from app.services.freesurfer_stats import StatsParseError, parse_stats_dir
from app.services.pipeline import Provenance

logger = get_logger(__name__)

# The values the Clinical Dementia Rating global score can take. Anything
# else in a CDR column is a data error, not a label.
VALID_CDR = (0.0, 0.5, 1.0, 2.0, 3.0)

# Spellings of "no value" seen in exported spreadsheets.
_MISSING_TOKENS = {"", "nan", "na", "n/a", "none", "null", "-", "?"}

# TODO(confirm): the cohort CSV's thickness column name was not specified
# ("thickness columns optional"). These spellings are accepted for a single
# mean-thickness column in mm; tell us the real header and we will pin it.
THICKNESS_COLUMNS = (
    "cortical_thickness_mm",
    "cortical_thickness",
    "mean_thickness_mm",
    "thickness_mm",
)

REQUIRED_CSV_COLUMNS = ("subject_id", "cdr", "age", "sex")


class CohortImportError(Exception):
    """The import cannot proceed at all (bad file, or --require-label unmet)."""


@dataclass(slots=True)
class ImportOptions:
    dry_run: bool = False
    # Fail the whole import, before writing anything, if any row lacks a CDR.
    require_label: bool = False
    include_mr2: bool = False
    source: DataSource = DataSource.OASIS
    scan_date: datetime.date | None = None
    # How the stats in a stats folder were produced.
    stats_extraction_method: ExtractionMethod = ExtractionMethod.FREESURFER_STATS


@dataclass(slots=True)
class ImportSummary:
    dry_run: bool = False
    imported: int = 0
    skipped_no_label: int = 0
    skipped_mr2: int = 0
    skipped_existing: int = 0
    skipped_invalid: int = 0
    qc_flagged: int = 0
    indeterminate: int = 0
    no_label_ids: list[str] = field(default_factory=list)
    # (subject id or row description, reason). Reasons never quote cell values.
    invalid: list[tuple[str, str]] = field(default_factory=list)

    def format(self) -> str:
        heading = "Cohort import summary"
        if self.dry_run:
            heading += " (DRY RUN: nothing was written)"
        lines = [
            heading,
            f"  imported:          {self.imported}",
            f"  skipped-no-label:  {self.skipped_no_label}",
            f"  skipped-MR2:       {self.skipped_mr2}",
            f"  skipped-existing:  {self.skipped_existing}",
            f"  skipped-invalid:   {self.skipped_invalid}",
            f"  QC-flagged:        {self.qc_flagged}",
            f"  indeterminate:     {self.indeterminate} of {self.imported} imported",
        ]
        for subject, reason in self.invalid:
            lines.append(f"  invalid  {subject}: {reason}")
        return "\n".join(lines)


@dataclass(frozen=True, slots=True)
class Candidate:
    """One session ready to import: validated, labelled, derived."""

    subject_id: str
    cdr: float
    age: int
    sex: Sex
    derived: DerivedBiomarkers
    extraction_method: ExtractionMethod


class _RowError(Exception):
    """This row cannot be imported; the message is safe to print."""


# --------------------------------------------------------------------------
# Cell parsing
# --------------------------------------------------------------------------


def is_missing(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return str(value).strip().lower() in _MISSING_TOKENS


def parse_cdr(value: Any) -> float | None:
    """The CDR in a cell, or None if the cell is empty.

    None means "not assessed". It is never replaced by a default: callers
    skip the row. Raises ValueError for a value that is not a CDR score.
    """
    if is_missing(value):
        return None
    try:
        cdr = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("cdr is not a number") from exc
    if cdr not in VALID_CDR:
        raise ValueError("cdr is not one of 0, 0.5, 1, 2, 3")
    return cdr


def _optional_float(row: dict[str, Any], column: str) -> float | None:
    value = row.get(column)
    if is_missing(value):
        return None
    try:
        number = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError) as exc:
        raise _RowError(f"{column} is not a number") from exc
    if not math.isfinite(number):
        raise _RowError(f"{column} is not finite")
    return number


def _positive_or_none(row: dict[str, Any], column: str) -> float | None:
    number = _optional_float(row, column)
    if number is not None and number <= 0:
        raise _RowError(f"{column} must be greater than zero")
    return number


def _parse_age(value: Any) -> int:
    if is_missing(value):
        raise _RowError("age is missing")
    try:
        age = float(value)
    except (TypeError, ValueError) as exc:
        raise _RowError("age is not a number") from exc
    if not 0 <= age <= 120:
        raise _RowError("age is outside 0-120")
    return int(round(age))


def _parse_sex(value: Any) -> Sex:
    if is_missing(value):
        raise _RowError("sex is missing")
    token = str(value).strip().upper()
    if token in ("M", "MALE"):
        return Sex.M
    if token in ("F", "FEMALE"):
        return Sex.F
    raise _RowError("sex is not M or F")


def _subject_id(value: Any, fallback: str) -> str:
    if is_missing(value):
        raise _RowError(f"{fallback}: subject_id is missing")
    try:
        subject_id = validate_external_subject_id(str(value))
    except ValueError as exc:
        raise _RowError(f"{fallback}: subject_id has an invalid shape") from exc
    assert subject_id is not None
    return subject_id


def is_mr2(subject_id: str) -> bool:
    """OASIS-1 names a subject's reliability rescan `<subject>_MR2`."""
    return subject_id.upper().endswith("_MR2")


# --------------------------------------------------------------------------
# Reading the inputs
# --------------------------------------------------------------------------


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.reader(handle)
            header = next(reader, None)
            if header is None:
                raise CohortImportError(f"{path.name} is empty.")
            headers = [h.strip() for h in header]
            return [dict(zip(headers, cells, strict=False)) for cells in reader if cells]
    except OSError as exc:
        raise CohortImportError(f"Could not read {path.name}.") from exc


def read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook  # imported here: only this path needs it

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any openpyxl/zip error is the same to the user
        raise CohortImportError(f"Could not read {path.name} as an .xlsx workbook.") from exc
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise CohortImportError(f"{path.name} is empty.")
        headers = [str(h).strip() if h is not None else "" for h in header]
        return [dict(zip(headers, row, strict=False)) for row in rows]
    finally:
        workbook.close()


def read_table(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return read_xlsx_rows(path)
    return read_csv_rows(path)


def _column(rows: list[dict[str, Any]], *names: str) -> str | None:
    """The header in `rows` matching one of `names`, case-insensitively."""
    if not rows:
        return None
    lowered = {str(h).strip().lower(): h for h in rows[0]}
    for name in names:
        if name.lower() in lowered:
            return lowered[name.lower()]
    return None


@dataclass(frozen=True, slots=True)
class Demographics:
    cdr_cell: Any
    age_cell: Any
    sex_cell: Any


def read_demographics(path: Path) -> dict[str, Demographics]:
    """The OASIS demographics sheet, keyed by subject id. Columns are found
    by name (`ID`, `M/F`, `Age`, `CDR`), in any order.
    """
    rows = read_table(path)
    id_col = _column(rows, "ID", "subject_id", "Subject ID")
    sex_col = _column(rows, "M/F", "sex", "gender")
    age_col = _column(rows, "Age")
    cdr_col = _column(rows, "CDR")
    absent = [
        name
        for name, col in (("ID", id_col), ("M/F", sex_col), ("Age", age_col), ("CDR", cdr_col))
        if col is None
    ]
    if absent:
        raise CohortImportError(f"{path.name} is missing the column(s) {', '.join(absent)}.")
    assert id_col and sex_col and age_col and cdr_col
    return {
        str(row[id_col]).strip(): Demographics(row.get(cdr_col), row.get(age_col), row.get(sex_col))
        for row in rows
        if not is_missing(row.get(id_col))
    }


# --------------------------------------------------------------------------
# Building candidates
# --------------------------------------------------------------------------


@dataclass(slots=True)
class _Screening:
    """Rows sorted into what will be imported and what will not, before any
    database write. `--require-label` is decided on this.
    """

    candidates: list[Candidate] = field(default_factory=list)
    summary: ImportSummary = field(default_factory=ImportSummary)


def _derived_from_csv_row(row: dict[str, Any]) -> DerivedBiomarkers:
    icv = _positive_or_none(row, "icv_mm3")
    hipp = _positive_or_none(row, "hipp_mm3")
    hipp_norm = _positive_or_none(row, "hipp_norm")
    vbr = _optional_float(row, "vbr")
    if vbr is not None and vbr < 0:
        raise _RowError("vbr must not be negative")
    cortex_gm_norm = _positive_or_none(row, "cortex_gm_norm")
    thickness = next(
        (t for t in (_positive_or_none(row, c) for c in THICKNESS_COLUMNS if c in row) if t), None
    )

    if hipp_norm is None and icv is not None and hipp is not None:
        hipp_norm = hipp / icv * 1000.0
    if hipp_norm is None and vbr is None:
        raise _RowError("neither hipp_norm nor vbr has a value")

    return derive_from_values(
        hippocampal_volume_norm=hipp_norm,
        ventricle_brain_ratio=vbr,
        cortical_thickness_mm=thickness,
        icv_mm3=icv,
        hippocampus_total_mm3=hipp,
        raw_volumes={
            "supplied_as": "cohort_csv",
            "icv_mm3": icv,
            "hippocampus_total_mm3": hipp,
            "cortical_gm_volume_norm": cortex_gm_norm,
        },
    )


def screen_csv(path: Path, options: ImportOptions) -> _Screening:
    rows = read_csv_rows(path)
    screening = _Screening(summary=ImportSummary(dry_run=options.dry_run))
    if not rows:
        return screening
    absent = [c for c in REQUIRED_CSV_COLUMNS if c not in rows[0]]
    if absent:
        raise CohortImportError(f"{path.name} is missing the column(s) {', '.join(absent)}.")

    for index, row in enumerate(rows, start=2):  # row 1 is the header
        label = f"row {index}"
        try:
            subject_id = _subject_id(row.get("subject_id"), label)
            label = subject_id
            if is_mr2(subject_id) and not options.include_mr2:
                screening.summary.skipped_mr2 += 1
                continue
            try:
                cdr = parse_cdr(row.get("cdr"))
            except ValueError as exc:
                raise _RowError(str(exc)) from exc
            if cdr is None:
                screening.summary.skipped_no_label += 1
                screening.summary.no_label_ids.append(subject_id)
                continue
            screening.candidates.append(
                Candidate(
                    subject_id=subject_id,
                    cdr=cdr,
                    age=_parse_age(row.get("age")),
                    sex=_parse_sex(row.get("sex")),
                    derived=_derived_from_csv_row(row),
                    extraction_method=ExtractionMethod.MANUAL_BIOMARKERS,
                )
            )
        except _RowError as exc:
            screening.summary.skipped_invalid += 1
            screening.summary.invalid.append((label, str(exc)))
    return screening


def _stats_dir_for(subject_dir: Path) -> Path:
    nested = subject_dir / "stats"
    return nested if nested.is_dir() else subject_dir


def screen_stats_folder(
    stats_root: Path, demographics_path: Path, options: ImportOptions
) -> _Screening:
    if not stats_root.is_dir():
        raise CohortImportError(f"{stats_root} is not a directory.")
    demographics = read_demographics(demographics_path)
    screening = _Screening(summary=ImportSummary(dry_run=options.dry_run))

    for subject_dir in sorted(p for p in stats_root.iterdir() if p.is_dir()):
        label = subject_dir.name
        try:
            subject_id = _subject_id(subject_dir.name, "directory")
            label = subject_id
            if is_mr2(subject_id) and not options.include_mr2:
                screening.summary.skipped_mr2 += 1
                continue
            person = demographics.get(subject_id)
            if person is None:
                raise _RowError("no row for this subject in the demographics sheet")
            try:
                cdr = parse_cdr(person.cdr_cell)
            except ValueError as exc:
                raise _RowError(str(exc)) from exc
            if cdr is None:
                screening.summary.skipped_no_label += 1
                screening.summary.no_label_ids.append(subject_id)
                continue
            try:
                derived = derive_from_stats(parse_stats_dir(_stats_dir_for(subject_dir)))
            except StatsParseError as exc:
                raise _RowError(exc.message) from exc
            screening.candidates.append(
                Candidate(
                    subject_id=subject_id,
                    cdr=cdr,
                    age=_parse_age(person.age_cell),
                    sex=_parse_sex(person.sex_cell),
                    derived=derived,
                    extraction_method=options.stats_extraction_method,
                )
            )
        except _RowError as exc:
            screening.summary.skipped_invalid += 1
            screening.summary.invalid.append((label, str(exc)))
    return screening


# --------------------------------------------------------------------------
# Writing
# --------------------------------------------------------------------------


async def import_screened(
    db: AsyncSession, screening: _Screening, options: ImportOptions
) -> ImportSummary:
    """Writes the screened candidates. Commits unless `options.dry_run`, in
    which case everything is rolled back after the summary is computed, so
    a dry run reports exactly what a real run would do.
    """
    summary = screening.summary
    if options.source == DataSource.PHANTOM:
        raise CohortImportError("A real cohort cannot be imported with source PHANTOM.")
    if options.require_label and summary.skipped_no_label:
        shown = ", ".join(summary.no_label_ids[:10])
        more = "" if summary.skipped_no_label <= 10 else ", ..."
        raise CohortImportError(
            f"--require-label: {summary.skipped_no_label} row(s) have no CDR ({shown}{more}). "
            "Nothing was imported."
        )

    scan_date = options.scan_date or datetime.date.today()
    seen: set[str] = set()
    try:
        for candidate in screening.candidates:
            existing = await patient_repo.get_scan_by_external_subject_id(db, candidate.subject_id)
            if existing is not None or candidate.subject_id in seen:
                summary.skipped_existing += 1
                continue
            seen.add(candidate.subject_id)

            patient = await patient_repo.create(
                db, age=candidate.age, sex=candidate.sex, source=options.source
            )
            scan = await patient_repo.create_scan(
                db,
                patient_id=patient.id,
                scan_date=scan_date,
                file_path=None,
                external_subject_id=candidate.subject_id,
                reference_label=candidate.cdr,
            )
            scan = await scan_service.process_derived_biomarkers(
                db,
                scan,
                candidate.derived,
                None,
                Provenance(
                    extraction_method=candidate.extraction_method,
                    data_domain=data_domain_for(options.source),
                    external_subject_id=candidate.subject_id,
                ),
            )
            summary.imported += 1
            if any(flag.get("kind") == "plausibility" for flag in scan.qc_flags or []):
                summary.qc_flagged += 1
            if scan.fis_result is not None and not scan.fis_result.determinate:
                summary.indeterminate += 1

        if options.dry_run:
            await db.rollback()
        else:
            await db.commit()
    except Exception:
        await db.rollback()
        raise

    logger.info(
        "cohort_import_finished",
        dry_run=options.dry_run,
        imported=summary.imported,
        skipped_no_label=summary.skipped_no_label,
        skipped_mr2=summary.skipped_mr2,
        skipped_existing=summary.skipped_existing,
        skipped_invalid=summary.skipped_invalid,
        qc_flagged=summary.qc_flagged,
    )
    return summary


async def import_csv(db: AsyncSession, path: Path, options: ImportOptions) -> ImportSummary:
    return await import_screened(db, screen_csv(path, options), options)


async def import_stats_folder(
    db: AsyncSession, stats_root: Path, demographics_path: Path, options: ImportOptions
) -> ImportSummary:
    return await import_screened(
        db, screen_stats_folder(stats_root, demographics_path, options), options
    )
